"""Prepare isolated browser-test fixtures without touching the development database."""

import os
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
(ROOT / "artifacts").mkdir(exist_ok=True)
os.environ["DATABASE_URL"] = "sqlite:///" + (ROOT / "artifacts" / "e2e.db").as_posix()
os.environ["SEMANTIC_ENABLED"] = "false"
sys.path.insert(0, str(ROOT / "backend"))

from app.db import Base, engine
from app.seed import seed

Base.metadata.create_all(engine)
seed()
document = Document()
document.add_heading("QA Candidate", 0)
document.add_paragraph(
    "Frontend developer building accessible React and TypeScript applications. Experienced with CSS and thoughtful product interfaces. Enjoys collaborating with cross-functional teams."
)
document.save(ROOT / "artifacts" / "qa-resume.docx")
print("Prepared isolated database and synthetic resume in artifacts/.")
