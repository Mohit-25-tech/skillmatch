from io import BytesIO
from unittest.mock import patch

from docx import Document
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import get_db
from app.main import app
from app.models import Job, Notification, Resume, Skill, WorkItem
from app.services.match_pipeline import match_resume
from app.services.matching import calculate_match
from app.services.product import create_alerts
from tests.conftest import account
from tests.test_api import JOB, resume_file


def upload(client, headers, content=None):
    return client.post(
        "/api/v1/resumes", headers=headers, files={"file": ("resume.docx", content or resume_file())}
    ).json()["id"]


def test_saved_state_match_filter_and_theme_preserves_scores(client):
    owner = account(client, "owner@example.com", "recruiter")
    job_id = client.post("/api/v1/jobs", headers=owner, json=JOB).json()["id"]
    candidate = account(client)
    assert client.get(f"/api/v1/jobs/{job_id}", headers=candidate).json()["saved"] is False
    client.post(f"/api/v1/jobs/{job_id}/save", headers=candidate)
    assert client.get(f"/api/v1/jobs/{job_id}", headers=candidate).json()["saved"] is True
    assert client.get("/api/v1/jobs?min_match=1", headers=candidate).json()["total"] == 0
    upload(client, candidate)
    score = client.get(f"/api/v1/jobs/{job_id}", headers=candidate).json()["score"]
    assert client.get(f"/api/v1/jobs?min_match={score}", headers=candidate).json()["total"] == 1
    assert client.get(f"/api/v1/jobs?min_match={score + 1}", headers=candidate).json()["total"] == 0
    from app.models import MatchResult

    with next(app.dependency_overrides[get_db]()) as db:
        before = list(db.scalars(select(MatchResult.id)))
        tasks = list(db.scalars(select(WorkItem.id)))
    profile = client.get("/api/v1/profile", headers=candidate).json()
    profile["preferences"]["theme"] = "light"
    response = client.put(
        "/api/v1/profile",
        headers=candidate,
        json={"name": profile["name"], "preferences": profile["preferences"]},
    )
    assert response.status_code == 200
    with next(app.dependency_overrides[get_db]()) as db:
        assert list(db.scalars(select(MatchResult.id))) == before
        assert list(db.scalars(select(WorkItem.id))) == tasks


def test_match_pages_do_not_serialize_the_entire_catalog(client):
    recruiter = account(client, "paging-owner@example.com", "recruiter")
    for number in range(14):
        response = client.post(
            "/api/v1/jobs",
            headers=recruiter,
            json={**JOB, "title": f"Python Engineer {number}"},
        )
        assert response.status_code == 201
    candidate = account(client)
    upload(client, candidate)
    first = client.get("/api/v1/matches?page=1&size=12", headers=candidate)
    second = client.get("/api/v1/matches?page=2&size=12", headers=candidate)
    assert first.status_code == second.status_code == 200
    assert len(first.json()) == 12 and len(second.json()) == 2
    assert {row["job_id"] for row in first.json()}.isdisjoint({row["job_id"] for row in second.json()})


def test_publisher_timestamps_normalize_to_utc():
    from app.ingestion.normalization import timestamp

    assert timestamp("2026-09-29T10:30:00+05:30") == timestamp("2026-09-29T05:00:00Z")
    assert timestamp("2026-09-29T05:00:00") == timestamp("2026-09-29T05:00:00Z")
    assert timestamp("unknown") is None


def test_upload_refreshes_dashboard_without_worker_or_job_browsing(client):
    recruiter = account(client, "poster@example.com", "recruiter")
    created = client.post("/api/v1/jobs", headers=recruiter, json=JOB)
    assert created.status_code == 201
    assert client.get("/api/v1/jobs").json()["items"][0]["id"] == created.json()["id"]
    candidate = account(client)
    upload(client, candidate)
    stats = client.get("/api/v1/analytics", headers=candidate).json()
    assert stats["matches"] == 1
    assert stats["average_score"] == 40.0
    doc = Document()
    doc.add_paragraph("Python FastAPI PostgreSQL developer")
    content = BytesIO()
    doc.save(content)
    upload(client, candidate, content.getvalue())
    stats = client.get("/api/v1/analytics", headers=candidate).json()
    assert stats["matches"] == 1
    assert stats["average_score"] == 60.0
    assert client.get("/api/v1/matches/status", headers=candidate).json()["pending_jobs"] == 0


def test_workspace_scores_latest_resume_and_tracker_privacy(client):
    recruiter = account(client, "owner@example.com", "recruiter")
    ids = [
        client.post("/api/v1/jobs", headers=recruiter, json={**JOB, "title": f"Python Engineer {n}"}).json()[
            "id"
        ]
        for n in range(3)
    ]
    candidate = account(client)
    saved = client.post(f"/api/v1/jobs/{ids[0]}/save", headers=candidate)
    assert saved.status_code == 201, saved.text
    assert client.get("/api/v1/applications", headers=recruiter).json() == []
    assert client.get("/api/v1/analytics", headers=candidate).json()["applications"] == 0
    resume_id = upload(client, candidate)
    jobs = client.get("/api/v1/jobs?size=2", headers=candidate).json()
    assert all(j["score"] == 40.0 for j in jobs["items"])
    rest = client.get("/api/v1/jobs", params={"cursor": jobs["next_cursor"]}, headers=candidate).json()
    assert len(rest["items"]) == 1
    assert len(client.get("/api/v1/matches", headers=candidate).json()) == 3
    applied = client.post(
        "/api/v1/applications", headers=candidate, json={"resume_id": resume_id, "job_id": ids[0]}
    )
    assert applied.status_code == 201, applied.text
    row = client.patch(
        f"/api/v1/tracker/{saved.json()['id']}",
        headers=candidate,
        json={"status": "Offer", "notes": "Follow up Friday"},
    ).json()
    assert [e["status"] for e in row["timeline"]] == ["Saved", "Applied", "Offer"]
    other = account(client, "other@example.com")
    assert (
        client.patch(f"/api/v1/tracker/{row['id']}", headers=other, json={"status": "Rejected"}).status_code
        == 404
    )
    doc = Document()
    doc.add_paragraph("Product manager with stakeholder communication and planning experience.")
    buffer = BytesIO()
    doc.save(buffer)
    latest = upload(client, candidate, buffer.getvalue())
    client.get("/api/v1/jobs", headers=candidate)
    matches = client.get("/api/v1/matches", headers=candidate).json()
    assert len(matches) == 3 and all(m["resume_id"] == latest for m in matches)
    assert client.get("/api/v1/analytics", headers=candidate).json()["matches"] == 3


def test_profile_resume_tools_and_public_discovery(client):
    owner = account(client, "owner@example.com", "recruiter")
    job_id = client.post(
        "/api/v1/jobs", headers=owner, json={**JOB, "salary_min": None, "salary_max": None}
    ).json()["id"]
    candidate = account(client)
    resume_id = upload(client, candidate)
    response = client.put(
        "/api/v1/profile",
        headers=candidate,
        json={
            "name": "Jane Smith",
            "preferences": {"theme": "light", "locations": ["Remote"], "remote_preference": "any"},
        },
    )
    assert response.status_code == 200
    assert client.get("/api/v1/profile", headers=candidate).json()["preferences"]["theme"] == "light"
    assert client.get(f"/api/v1/resumes/{resume_id}/ats", headers=candidate).json()["word_count"] > 0
    tailoring = client.get(f"/api/v1/jobs/{job_id}/tailoring", headers=candidate)
    assert tailoring.status_code == 200, tailoring.text
    assert tailoring.json()["match"]["improvements"]
    assert client.post(f"/api/v1/jobs/{job_id}/draft", headers=candidate).status_code in (200, 503)
    assert client.get("/api/v1/search?q=Python").json()["jobs"]
    assert client.get("/api/v1/insights/market").json()["salaries"] == []
    assert client.get("/api/v1/learning-path", headers=candidate).json()["items"][0]["related_jobs"] == 1
    assert client.get(f"/api/v1/jobs/{job_id}/similar").status_code == 200
    # Non-applicants are excluded without discoverability consent.
    assert client.get(f"/api/v1/jobs/{job_id}/candidates", headers=owner).json()["items"] == []
    client.put(
        "/api/v1/profile",
        headers=candidate,
        json={"name": "Jane Smith", "preferences": {"discoverable": True}},
    )
    client.get("/api/v1/jobs", headers=candidate)
    ranked = client.get(f"/api/v1/jobs/{job_id}/candidates", headers=owner).json()["items"]
    assert set(ranked[0]) == {"candidate_id", "name", "score"}


def test_saved_search_notifications_and_ownership(client):
    candidate = account(client)
    search = client.post(
        "/api/v1/saved-searches",
        headers=candidate,
        json={"name": "Python", "filters": {"keywords": "Python", "min_match": 20}},
    )
    assert search.status_code == 201
    owner = account(client, "owner@example.com", "recruiter")
    client.post("/api/v1/jobs", headers=owner, json=JOB)
    upload(client, candidate)
    override = app.dependency_overrides[get_db]
    with next(override()) as db:
        resume = db.scalar(select(Resume))
        match_resume(db, resume)
        create_alerts(db, resume)
        db.commit()
        assert len(db.scalars(select(Notification)).all()) == 1
    notifications = client.get("/api/v1/notifications", headers=candidate).json()
    assert notifications["unread"] == 1
    notification_id = notifications["items"][0]["id"]
    other = account(client, "other@example.com")
    assert client.post(f"/api/v1/notifications/{notification_id}/read", headers=other).status_code == 404
    assert client.delete(f"/api/v1/saved-searches/{search.json()['id']}", headers=other).status_code == 404
    assert client.post(f"/api/v1/notifications/{notification_id}/read", headers=candidate).status_code == 200
    assert client.get("/api/v1/notifications", headers=candidate).json()["unread"] == 0
    assert client.get("/api/v1/notifications/stream").status_code == 401


def test_weighted_experience_location_and_optional_skills():
    python, sql = Skill(name="Python"), Skill(name="SQL")
    resume = Resume(text="3 years of experience", skills=[python], experience_years=3)
    job = Job(
        title="Engineer",
        description="5 years experience",
        skills=[python, sql],
        skill_importance={"SQL": "nice-to-have"},
        experience_min=5,
        location="London",
        remote=False,
    )
    result = calculate_match(resume, job, {"locations": ["London"]})
    assert result["keyword_score"] == 29.4
    assert result["components"]["experience"]["score"] == 60
    assert result["components"]["location"]["score"] == 100
    assert result["components"]["semantic"]["score"] is None
    assert result["score"] == 45.6
    assert result["confidence"] == "low"
    assert result["confidence_label"] == "Low confidence"


def test_worker_durable_resume_task(client):
    owner = account(client, "owner@example.com", "recruiter")
    for n in range(4):
        client.post("/api/v1/jobs", headers=owner, json={**JOB, "title": f"Python Engineer {n}"})
    candidate = account(client)
    resume_id = upload(client, candidate)
    with next(app.dependency_overrides[get_db]()) as db:
        bind = db.get_bind()
        # Prior native catalog tasks are independently exercised by all-task processing.
    from app.worker import process_tasks

    with patch("app.worker.SessionLocal", side_effect=lambda: Session(bind)):
        for _ in range(5):
            process_tasks()
    status = client.get("/api/v1/matches/status", headers=candidate).json()
    assert status["resume_id"] == resume_id and status["scored_jobs"] == 4 and status["pending_jobs"] == 0
    with Session(bind) as db:
        assert set(db.scalars(select(WorkItem.status))) == {"done"}


def test_admin_ingestion_controls_and_digest_opt_in(client):
    import smtplib

    from app.config import get_settings
    from app.models import SavedSearch, User
    from app.services.product import send_digests

    admin = account(client, "admin@example.com")
    with next(app.dependency_overrides[get_db]()) as db:
        db.scalar(select(User).where(User.email == "admin@example.com")).role = "admin"
        db.commit()
    created = client.post(
        "/api/v1/admin/ingestion",
        headers=admin,
        json={"key": "test", "kind": "greenhouse", "config": {"slug": "acme"}},
    )
    assert created.status_code == 201
    source_id = created.json()["id"]
    assert client.post(f"/api/v1/admin/ingestion/{source_id}/run", headers=admin).status_code == 409
    assert (
        client.patch(
            f"/api/v1/admin/ingestion/{source_id}",
            headers=admin,
            json={"enabled": True, "config": {"slug": "new-acme"}},
        ).status_code
        == 200
    )
    first = client.post(f"/api/v1/admin/ingestion/{source_id}/run", headers=admin).json()
    assert (
        client.post(f"/api/v1/admin/ingestion/{source_id}/run", headers=admin).json()["work_item_id"]
        == first["work_item_id"]
    )
    candidate = account(client)
    assert client.get("/api/v1/admin/ingestion", headers=candidate).status_code == 403
    owner = account(client, "owner@example.com", "recruiter")
    job_id = client.post("/api/v1/jobs", headers=owner, json=JOB).json()["id"]
    with next(app.dependency_overrides[get_db]()) as db:
        user = db.scalar(select(User).where(User.email == "candidate@example.com"))
        search = SavedSearch(user_id=user.id, name="All", filters={})
        db.add(search)
        db.flush()
        notification = Notification(user_id=user.id, search_id=search.id, job_id=job_id, title="New job")
        db.add(notification)
        db.commit()
        with (
            patch.object(get_settings(), "smtp_host", "mailpit"),
            patch("app.services.product.smtplib.SMTP") as smtp,
        ):
            assert send_digests(db) == 0
            smtp.assert_not_called()
            user.preferences = {"email_digest": True}
            db.commit()
            smtp.return_value.__enter__.return_value.send_message.side_effect = smtplib.SMTPException(
                "offline"
            )
            import pytest

            with pytest.raises(smtplib.SMTPException):
                send_digests(db)
            assert notification.emailed_at is None
            smtp.return_value.__enter__.return_value.send_message.side_effect = None
            assert send_digests(db) == 1
            assert notification.emailed_at is not None
            assert send_digests(db) == 0
