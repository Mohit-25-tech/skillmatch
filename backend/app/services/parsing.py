from functools import lru_cache
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pdfplumber
import spacy
from docx import Document
from fastapi import HTTPException
from spacy.matcher import PhraseMatcher

from app.services.taxonomy import ALIASES

MAX_FILE_SIZE = 5 * 1024 * 1024


def parse_resume(data: bytes, filename: str) -> str:
    if not data or len(data) > MAX_FILE_SIZE:
        raise HTTPException(413, "Upload a non-empty file smaller than 5 MB")
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".pdf" and data.startswith(b"%PDF-"):
            with pdfplumber.open(BytesIO(data)) as pdf:
                if len(pdf.pages) > 20:
                    raise HTTPException(422, "Please use a resume with 20 pages or fewer")
                text = "\n".join((page.extract_text() or "") for page in pdf.pages)
        elif extension == ".docx" and data.startswith(b"PK"):
            with ZipFile(BytesIO(data)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 25 * 1024 * 1024:
                    raise HTTPException(413, "The expanded document is too large")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("Invalid DOCX")
            doc = Document(BytesIO(data))
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [c.text for t in doc.tables for r in t.rows for c in r.cells]
            )
        else:
            raise HTTPException(415, "Only genuine PDF and DOCX files are supported")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Unable to read this file. Use an unencrypted PDF or DOCX.") from exc
    if len(text.strip()) < 30:
        raise HTTPException(422, "No readable resume text found. Scanned PDFs need OCR before upload.")
    return text[:100000]


@lru_cache(maxsize=8)
def skill_matcher(names: tuple[str, ...]):
    nlp = spacy.blank("en")
    matcher = PhraseMatcher(nlp.vocab, attr="LOWER")
    for name in names:
        matcher.add(name, [nlp.make_doc(name)])
    canonical_names = {name.casefold(): name for name in names}
    for alias, canonical in ALIASES.items():
        if canonical.casefold() in canonical_names:
            matcher.add(canonical_names[canonical.casefold()], [nlp.make_doc(alias)])
    return nlp, matcher


import re


def generate_parse_warnings(text: str, skills: list[str]) -> list[str]:
    warnings = []
    # Contact check
    has_email = bool(re.search(r"[\w.+-]+@[\w.-]+\.[a-z]{2,}", text, re.I))
    has_phone = bool(re.search(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", text))
    if not has_email and not has_phone:
        warnings.append("No contact email or phone number detected. Recruiters may have difficulty reaching you.")
    elif not has_email:
        warnings.append("No email address detected. Adding a direct email address is recommended.")

    # Skills check
    if len(skills) < 3:
        warnings.append(
            f"Only {len(skills)} skill(s) detected. Listing technical tools in a dedicated Skills section improves ATS match scores."
        )

    # Experience or projects check
    has_exp = bool(re.search(r"\b(experience|employment|work history|projects|internships|academic projects)\b", text, re.I))
    if not has_exp:
        warnings.append("No Experience, Projects, or Internships section detected. Adding coursework projects or past roles enhances ATS evaluation.")

    # Education check
    has_edu = bool(re.search(r"\b(education|degree|university|college|bachelor|master|b\.?tech|b\.?e|b\.?sc|m\.?tech|mca|bca)\b", text, re.I))
    if not has_edu:
        warnings.append("No Education section detected. Consider listing your degree, college, or graduation year.")

    # Length check
    words = len(text.split())
    if words < 120:
        warnings.append("Resume content appears very short (<120 words). Ensure sections were not clipped.")

    # Artifacts check
    if "\ufffd" in text:
        warnings.append("Encoding artifacts found in document text. Consider re-exporting your resume cleanly.")

    return warnings


def extract_skills(text: str, names: list[str]) -> list[str]:
    nlp, matcher = skill_matcher(tuple(sorted(set(names))))
    return sorted({nlp.vocab.strings[match_id] for match_id, _, _ in matcher(nlp.make_doc(text))})
